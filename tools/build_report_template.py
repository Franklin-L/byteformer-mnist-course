"""Convert the student Markdown worksheet to an editable Word handout."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def plain(text):
    return text.replace('**', '').replace('`', '').strip()


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(11 if name == 'Normal' else 20 if name == 'Title' else 14)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.core_properties.title = '码流图像分类实验报告模板'
    document.core_properties.author = 'Franklin-L'
    lines = (ROOT/'docs/student_report_template.md').read_text().splitlines()
    i, in_code = 0, False
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            in_code = not in_code; i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [plain(c) for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c) for c in cells): rows.append(cells)
                i += 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Table Grid'
            for j,c in enumerate(rows[0]):
                table.rows[0].cells[j].text = c
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'DBEAF4')
                table.rows[0].cells[j]._tc.get_or_add_tcPr().append(shading)
            for row in rows[1:]:
                for cell, value in zip(table.add_row().cells,row): cell.text = value or ' '
            document.add_paragraph()
            continue
        if line.startswith('# '): document.add_heading(plain(line[2:]), 0)
        elif line.startswith('## '): document.add_heading(plain(line[3:]), 1)
        elif line.startswith('### '): document.add_heading(plain(line[4:]), 2)
        elif line.strip():
            text=plain(line.lstrip('> '))
            if text.startswith('- [ ]'): text='□ ' + text[5:].strip()
            elif text.startswith('- '): text='• ' + text[2:]
            paragraph=document.add_paragraph(text)
            if in_code:
                for run in paragraph.runs: run.font.color.rgb=RGBColor.from_string('176BA0')
        i += 1
    footer=section.footer.paragraphs[0]
    footer.text='码流图像分类实验报告'
    document.save(ROOT/'docs/学生实验报告模板.docx')
    print(ROOT/'docs/学生实验报告模板.docx')


if __name__ == '__main__': main()
